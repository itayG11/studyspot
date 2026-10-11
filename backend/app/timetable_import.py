"""Reading a timetable exported from the college ("Yedion", an Excel file):
a preview for the student to check, never saved as is, and the file itself
is never kept.

The export is a grid: a row with the days, a column of hours, and in the
cells the classes, such as
    61739-מבני נתונים תרגיל<TAB>גב' ... 219 M נגישות M 08:30-10:30
and "המשך ..." where a class goes on into the next hour. From each class
only the day, the times, the building and the course name are read: the
lecturer's name and the course number are not kept.

The file comes from a stranger's computer, so before any parsing: a small
size, a real zip (an .xlsx is one), and a small size once unpacked (a zip
bomb is tiny on the wire and huge inside). openpyxl reads its XML through
defusedxml (a dependency), which refuses the XML tricks (entity bombs,
external entities)."""

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass

from openpyxl import load_workbook

from app.models import Building

MAX_BYTES = 200 * 1024
MAX_UNPACKED = 5 * 1024 * 1024
MAX_ROWS = 40  # app/classes.py MAX_CLASSES
TITLE_LENGTH = 60
PLACE_LENGTH = 40
# Only the corner a timetable fills: a cell far out (column XFD) would
# otherwise turn every row into thousands of empty cells in memory.
READ_ROWS, READ_COLUMNS = 200, 12
# A real class is a line or two; a longer cell is not one (and is not read).
MAX_CELL = 400

WEEKDAYS = {"ראשון": 6, "שני": 0, "שלישי": 1, "רביעי": 2, "חמישי": 3, "שישי": 4, "שבת": 5}
# The times at the end of a class: "08:30-10:30", also with an en dash.
_TIMES = re.compile(r"(\d{1,2}):(\d{2}) ?[-–] ?(\d{1,2}):(\d{2})$")
_NUMBER = re.compile(r"\d+-")  # the course number, before its name
KINDS = ("הרצאה", "תרגיל", "תרגול", "מעבדה", "סמינר")


class ImportRefused(Exception):
    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class Row:
    weekday: int
    starts: str  # "HH:MM"
    ends: str
    building_code: str | None  # None: no building of the campus by that name
    building_text: str  # what the file says, to show next to the choice
    title: str | None


def read_yedion(data: bytes, buildings: list[Building]) -> list[Row]:
    sheet = _open(data)
    rows = list(sheet.iter_rows(max_row=READ_ROWS, max_col=READ_COLUMNS, values_only=True))
    days, start = _day_columns(rows)
    out: list[Row] = []
    for values in rows[start:]:
        for column, weekday in days.items():
            for row in _read_cell(values[column], weekday, buildings) if column < len(values) else []:
                if row not in out:
                    out.append(row)
    if not out:
        raise ImportRefused("import_empty")
    if len(out) > MAX_ROWS:
        raise ImportRefused("import_too_many")
    return sorted(out, key=lambda r: ((r.weekday + 1) % 7, r.starts))


def _open(data: bytes):
    if len(data) > MAX_BYTES:
        raise ImportRefused("import_too_big")
    if not data.startswith(b"PK\x03\x04"):
        raise ImportRefused("import_not_xlsx")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNPACKED:
                raise ImportRefused("import_too_big")
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except ImportRefused:
        raise
    except Exception:  # noqa: BLE001 (a broken zip, not a workbook, bad XML: all one answer to the student)
        raise ImportRefused("import_not_xlsx") from None
    return book.worksheets[0]


def _day_columns(rows: list[tuple]) -> tuple[dict[int, int], int]:
    """The row naming the days ("יום ראשון" ...): column -> weekday, and the
    first row after it."""
    for index, values in enumerate(rows[:10]):
        days = {}
        for column, value in enumerate(values):
            name = str(value or "").replace("יום", "").strip()
            if name in WEEKDAYS:
                days[column] = WEEKDAYS[name]
        if len(days) >= 2:
            return days, index + 1
    raise ImportRefused("import_not_yedion")


def _read_cell(value, weekday: int, buildings: list[Building]) -> list[Row]:
    """A cell can hold more than one class, a line each."""
    text = str(value or "")
    if len(text) > MAX_CELL:
        return []
    out = []
    for line in _clean(text).splitlines():
        line = line.strip()
        if line and not line.startswith("המשך") and (row := _read_class(line, weekday, buildings)) is not None:
            out.append(row)
    return out


def _read_class(line: str, weekday: int, buildings: list[Building]) -> Row | None:
    times = _TIMES.search(" ".join(line.split()))
    if times is None:
        return None
    sh, sm, eh, em = (int(g) for g in times.groups())
    if not (0 <= sh < 24 and 0 <= eh < 24 and sm < 60 and em < 60) or (sh, sm) >= (eh, em):
        return None
    words = " ".join(line.split()).split(" ")
    place = words[: -len(times.group(0).split(" "))]
    code, said = _building(place, buildings)
    return Row(weekday, f"{sh:02d}:{sm:02d}", f"{eh:02d}:{em:02d}", code, said, _title(line))


def _title(line: str) -> str | None:
    """"מבני נתונים, תרגיל": the course and the kind of class, nothing about
    who teaches it. Words, not a pattern: no input can make this slow."""
    first = line.split("\t")[0].split()
    if not first or not _NUMBER.match(first[0]):
        return None
    words = [_NUMBER.sub("", first[0], count=1), *first[1:]]
    for index, word in enumerate(words):
        if word in KINDS and index > 0:
            name = " ".join(w for w in words[:index] if w)
            return f"{name}, {word}"[:TITLE_LENGTH] if name else None
    return None


def _building(place: list[str], buildings: list[Building]) -> tuple[str | None, str]:
    """Where the class is, from the words just before its times only (never
    further back: that is the lecturer's name). The export writes the place
    twice: "219 M נגישות M", "702 L L", "אולם ספורט אולם ספורט"."""
    codes = {b.code.casefold(): b.code for b in buildings}
    for word in reversed(place[-2:]):
        if word.casefold() in codes:
            return codes[word.casefold()], word
    said = _twice(place)[:PLACE_LENGTH]  # app/schemas.py ClassSlotIn.place_text
    for b in buildings:
        if said and b.name and b.name == said:
            return b.code, said
    return None, said


def _twice(words: list[str]) -> str:
    """The words written twice at the end ("אולם ספורט אולם ספורט" ->
    "אולם ספורט"), or nothing: anything else could be a person's name."""
    for size in range(min(4, len(words) // 2), 0, -1):
        if words[-size:] == words[-2 * size : -size]:
            return " ".join(words[-size:])
    return ""


def _clean(text: str) -> str:
    """No control or direction characters (they could hide or reorder text);
    tabs stay, they split the course from the rest, and line breaks, which
    split one class from the next."""
    return "".join(ch for ch in text if ch in "\t\n" or unicodedata.category(ch) not in ("Cc", "Cf")).strip()
