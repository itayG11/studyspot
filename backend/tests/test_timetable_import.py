"""Importing a timetable from the college's export (the "Yedion" Excel
file): a preview only, from the file's text, never kept. Built here in code
with the same layout as a real export, with made-up lecturers."""

import io
import time
import zipfile

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select

from app.models import ClassSlot
from app.timetable_import import MAX_BYTES, ImportRefused, read_yedion, without_titles

DAYS = ["יום ראשון", "יום שני", "יום שלישי", "יום רביעי", "יום חמישי", "יום שישי"]
# Column by day, row by hour, as the export has them.
CELLS = {
    ("יום ראשון", "08:30"): "61739-מבני נתונים ומבוא לאלגוריתמים תרגיל\tגב' כהן דנה 219 M נגישות M 08:30-10:30",
    ("יום ראשון", "09:30"): "המשך\tמבני נתונים ומבוא לאלגוריתמים",
    ("יום ראשון", "10:30"): '61911-מבנים אלגבריים להנדסת תוכנה 2 הרצאה\tד"ר לוי רון 317 M M 10:30-12:20',
    ("יום ראשון", "12:50"): "11129-טורים, התמרות ומשוואות דיפרנציאליות הרצאה\tפרופ' אברהם יוסי 702 L L 12:50-15:50",
    ("יום שני", "08:30"): "61739-מבני נתונים ומבוא לאלגוריתמים הרצאה גב' כהן דנה\t209 M M 08:30-11:30",
    ("יום חמישי", "10:30"): "11575-מובילי אורח חיים בריא הרצאה מר שמש גל\tאולם ספורט אולם ספורט 10:30-12:20",
}
HOURS = ["08:30", "09:30", "10:30", "11:30", "12:50"]


def yedion(cells=CELLS, title="טבלה - לוח מערכת שעות") -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append([title])
    sheet.append([""] + DAYS)
    for hour in HOURS:
        sheet.append([hour] + [cells.get((day, hour), "") for day in DAYS])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_reads_each_class_with_its_day_times_building_and_course(session, braude):
    rows = read_yedion(yedion(), braude.buildings)
    got = sorted((r.weekday, r.starts, r.ends, r.building_code, r.title) for r in rows)
    assert got == sorted([
        (6, "08:30", "10:30", "M", "מבני נתונים ומבוא לאלגוריתמים, תרגיל"),
        (6, "10:30", "12:20", "M", "מבנים אלגבריים להנדסת תוכנה 2, הרצאה"),
        (6, "12:50", "15:50", "L", "טורים, התמרות ומשוואות דיפרנציאליות, הרצאה"),
        (0, "08:30", "11:30", "M", "מבני נתונים ומבוא לאלגוריתמים, הרצאה"),
        (3, "10:30", "12:20", None, "מובילי אורח חיים בריא, הרצאה"),
    ])


def test_a_building_not_on_campus_keeps_what_the_file_says(session, braude):
    # The export writes the place twice ("אולם ספורט אולם ספורט"): that, whole.
    gym = next(r for r in read_yedion(yedion(), braude.buildings) if r.building_code is None)
    assert gym.building_text == "אולם ספורט"


def test_lecturers_and_course_numbers_are_not_read(session, braude):
    text = " ".join(f"{r.title} {r.building_text}" for r in read_yedion(yedion(), braude.buildings))
    for kept_out in ("כהן", "לוי", "אברהם", "שמש", "61739", "גב'"):
        assert kept_out not in text


def test_the_preview_saves_nothing(client, session):
    assert client.post("/me/classes/import", content=yedion()).status_code == 200
    assert session.scalar(select(func.count()).select_from(ClassSlot)) == 0


def cell(text):
    return yedion(cells={("יום ראשון", "08:30"): text})


def test_two_classes_in_one_cell_are_two_classes(session, braude):
    both = "1-אלגברה תרגיל\tמר א 101 M M 08:30-09:30\n2-חדו\"א הרצאה\tמר ב 202 L L 10:30-12:30"
    rows = read_yedion(cell(both), braude.buildings)
    assert [(r.starts, r.building_code, r.title) for r in rows] == [
        ("08:30", "M", "אלגברה, תרגיל"),
        ("10:30", "L", 'חדו"א, הרצאה'),
    ]


def test_a_lecturer_named_like_a_building_is_not_the_building(session, braude):
    [row] = read_yedion(cell("11575-ספורט הרצאה מר L כהן\tאולם ספורט אולם ספורט 10:30-12:20"), braude.buildings)
    assert (row.building_code, row.building_text) == (None, "אולם ספורט")


def test_a_cell_with_no_room_says_nothing_of_the_lecturer(session, braude):
    [row] = read_yedion(cell("61774-יסודות המחשוב הרצאה\tמר ישראל ישראלי 10:30-12:20"), braude.buildings)
    assert (row.building_code, row.building_text) == (None, "")


def test_times_with_an_en_dash_are_read(session, braude):
    [row] = read_yedion(cell("1-אלגברה תרגיל\t101 M M 8:30–10:00"), braude.buildings)
    assert (row.starts, row.ends) == ("08:30", "10:00")


def test_a_cell_built_to_slow_the_reader_is_skipped_quickly(session, braude):
    # Many spaces between a course and its times once took the reader many
    # seconds (catastrophic backtracking); such a cell is not a class.
    started = time.monotonic()
    with pytest.raises(ImportRefused):
        read_yedion(cell("1-a" + " " * 100_000 + "x 08:30-10:30"), braude.buildings)
    assert time.monotonic() - started < 2


def test_a_very_wide_sheet_is_read_in_its_corner_only(session, braude):
    # A cell far to the left (column XFD) made every row a row of 16,384 cells.
    out = io.BytesIO()
    book = Workbook()
    sheet = book.active
    sheet.append([""] + DAYS)
    sheet.append(["08:30", "1-אלגברה תרגיל\t101 M M 08:30-10:00"])
    for row in range(3, 400):
        sheet.cell(row=row, column=16384, value="x")
    book.save(out)
    rows = read_yedion(out.getvalue(), braude.buildings)
    assert [r.title for r in rows] == ["אלגברה, תרגיל"]


def test_the_xml_is_read_through_defusedxml():
    import openpyxl.xml

    assert openpyxl.xml.DEFUSEDXML and not openpyxl.xml.LXML


@pytest.mark.parametrize(
    ("data", "code"),
    [
        (b"x" * (MAX_BYTES + 1), "import_too_big"),
        (b"%PDF-1.7 not a spreadsheet", "import_not_xlsx"),
        (b"PK\x03\x04 broken zip", "import_not_xlsx"),
    ],
)
def test_refuses_what_is_not_a_small_excel_file(session, braude, data, code):
    with pytest.raises(ImportRefused) as refused:
        read_yedion(data, braude.buildings)
    assert refused.value.code == code


def test_refuses_a_zip_bomb_before_opening_it(session, braude):
    # Small on the wire, huge unpacked: the sizes inside are checked first.
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/worksheets/sheet1.xml", b"0" * (20 * 1024 * 1024))
    with pytest.raises(ImportRefused) as refused:
        read_yedion(out.getvalue(), braude.buildings)
    assert refused.value.code == "import_too_big"


def test_an_excel_file_that_is_not_a_timetable_is_refused(session, braude):
    book = Workbook()
    book.active.append(["שם", "ציון"])
    out = io.BytesIO()
    book.save(out)
    with pytest.raises(ImportRefused) as refused:
        read_yedion(out.getvalue(), braude.buildings)
    assert refused.value.code == "import_not_yedion"


def test_a_timetable_with_no_classes_is_refused(session, braude):
    with pytest.raises(ImportRefused) as refused:
        read_yedion(yedion(cells={}), braude.buildings)
    assert refused.value.code == "import_empty"


def test_the_api_previews_for_the_signed_in_student(client):
    response = client.post(
        "/me/classes/import", content=yedion(), headers={"Content-Type": "application/octet-stream"}
    )
    assert response.status_code == 200
    body = response.json()
    assert len(body) == 5 and {"weekday", "starts", "ends", "building_code", "building_text", "title"} <= set(body[0])


def test_the_api_needs_a_signed_in_student(client):
    client.user = None
    assert client.post("/me/classes/import", content=yedion()).status_code == 401


def test_the_api_says_why_a_file_is_refused(client):
    response = client.post("/me/classes/import", content=b"not a file")
    assert (response.status_code, response.json()["detail"]) == (422, "import_not_xlsx")


def test_adding_many_at_once_says_which_were_refused(client, session):
    body = [
        {"weekday": 6, "starts": "08:30", "ends": "10:00", "building_code": "M", "title": "א, הרצאה"},
        {"weekday": 6, "starts": "09:00", "ends": "11:00", "building_code": "M", "title": "ב, תרגיל"},  # overlaps
        {"weekday": 0, "starts": "08:30", "ends": "10:00", "building_code": "NOPE", "title": None},
    ]
    response = client.post("/me/classes/bulk", json=body)
    assert response.status_code == 200
    assert [r["refused"] for r in response.json()] == [None, "class_overlaps", "building_not_found"]
    assert session.scalar(select(func.count()).select_from(ClassSlot)) == 1


def test_adding_many_at_once_is_capped(client):
    one = {"weekday": 6, "starts": "08:30", "ends": "09:00", "building_code": "M", "title": None}
    assert client.post("/me/classes/bulk", json=[one] * 41).status_code == 422


def test_a_lecturers_title_is_not_part_of_the_course(session, braude):
    # From the user's check, 25 Oct: "ד"ר" had come into the course's name.
    [row] = read_yedion(cell('61739-ד"ר מבנים אלגבריים 2 להנדסת תוכנה תרגיל\tמר א 101 M M 08:30-09:30'), braude.buildings)
    assert row.title == "מבנים אלגבריים 2 להנדסת תוכנה, תרגיל"


@pytest.mark.parametrize(
    "name",
    ['ד"ר מבנים אלגבריים', "ד״ר מבנים אלגבריים", "ד''ר מבנים אלגבריים", "פרופ׳ מבנים אלגבריים", "מבנים אלגבריים גב'", "מבנים, ד”ר אלגבריים"],
)
def test_titles_are_taken_out_however_the_quotes_are_written(name):
    assert without_titles(name).replace(",", "") == "מבנים אלגבריים"


def test_a_word_that_only_contains_a_title_stays():
    # Whole words only: "מרכזי" holds "מר", "דרך" holds "דר".
    assert without_titles("מבוא לתכנון מרכזי דרך") == "מבוא לתכנון מרכזי דרך"
