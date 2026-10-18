"""An institution admin uploads a table of buildings and places (Excel or
CSV): a preview first, nothing saved, then the whole table in one go. A
building that is already there is the admin's call: add the new places to
it, leave it as it is (not sent), or delete it and add it again from the
file."""

import io
import zipfile
from datetime import UTC, datetime, time, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app import campus_import
from app.config import Settings, get_settings
from app.main import app
from app.models import (
    Booking,
    BookingSource,
    BookingStatus,
    Building,
    BuildingStatus,
    ClassSlot,
    Place,
    Seat,
    User,
    UserRole,
)
from app.seed import seed_demo

HEADER = "שם קצר,שם הבניין,מספר קומות,סוג המקום,שם המקום,קומה,כמה אנשים,שורות,עמדות בשורה"


@pytest.fixture
def admin(session, braude) -> User:
    user = User(institution_id=braude.id, email="a@braude.ac.il", display_name="A", role=UserRole.INSTITUTION_ADMIN)
    session.add(user)
    session.flush()
    return user


@pytest.fixture
def demo(session):
    return seed_demo(session)


@pytest.fixture
def demo_admin(session, demo) -> User:
    user = User(institution_id=demo.id, email="d@x.example", display_name="D", role=UserRole.INSTITUTION_ADMIN)
    session.add(user)
    session.flush()
    return user


@pytest.fixture
def demo_on():
    settings = Settings(database_url="postgresql+psycopg://unused@localhost/unused", demo_login_enabled=True, _env_file=None)
    app.dependency_overrides[get_settings] = lambda: settings
    yield settings
    app.dependency_overrides.pop(get_settings, None)


@pytest.fixture(autouse=True)
def committed_seed(session, braude):
    """A refused import rolls back; the seed data must be there before it,
    as in production, not only in the test's open transaction."""
    session.commit()


def csv(*lines: str, header: str = HEADER, encoding: str = "utf-8") -> bytes:
    return "\n".join([header, *lines]).encode(encoding)


def xlsx(rows: list[list]) -> bytes:
    book = Workbook()
    for row in rows:
        book.active.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def preview(client, data: bytes, slug: str = "braude"):
    return client.post(
        f"/admin/institutions/{slug}/import/preview", content=data,
        headers={"content-type": "application/octet-stream"},
    )


def save(client, buildings: list[dict], slug: str = "braude"):
    return client.post(f"/admin/institutions/{slug}/import", json={"buildings": buildings})


def to_save(previewed: dict, action: str = "add", **extra) -> list[dict]:
    """What the site sends after the preview: every building, its places."""
    keep = ("kind", "name", "floor", "capacity", "lab_rows", "lab_cols", "location_note", "weekday_hours", "friday_hours")
    out = []
    for b in previewed["buildings"]:
        item = {k: b[k] for k in ("code", "name", "floors_count", "status", "latitude", "longitude")}
        item["places"] = [{k: p[k] for k in keep} for p in b["places"] if p["when_added"] is None or action == "replace"]
        if b["existing"] is not None:
            item["action"] = action
            if action == "replace":
                item["expected_places"] = b["existing"]["places"]
                item["expected_upcoming_bookings"] = b["existing"]["upcoming_bookings"]
        out.append({**item, **extra})
    return out


def count(session, model, **where) -> int:
    query = select(func.count()).select_from(model)
    for key, value in where.items():
        query = query.where(getattr(model, key) == value)
    return session.scalar(query)


# --- Reading the file ------------------------------------------------------------------


def test_a_csv_previews_new_buildings_and_places_and_saves_nothing(client, admin, session):
    client.user = admin
    data = csv(
        "T,בניין הניסיון,2,חדר לימוד קבוצתי,T101,0,8,,",
        "T,בניין הניסיון,2,חוות מחשבים,T201,1,,3,4",
        "Q,,1,,,,,,",
    )
    response = preview(client, data)
    assert response.status_code == 200, response.json()
    body = response.json()
    assert body["problems"] == []
    t, q = body["buildings"]
    assert (t["code"], t["name"], t["floors_count"], t["existing"]) == ("T", "בניין הניסיון", 2, None)
    assert [(p["line"], p["kind"], p["name"], p["capacity"], p["lab_rows"]) for p in t["places"]] == [
        (2, "group_room", "T101", 8, None), (3, "computer_lab", "T201", None, 3),
    ]
    assert (q["code"], q["places"]) == ("Q", [])
    assert count(session, Building, code="T") == 0


def test_the_whole_table_is_saved_with_signs_hours_and_lab_stations(client, admin, session, braude):
    client.user = admin
    body = preview(client, csv("T,בניין הניסיון,2,חוות מחשבים,T201,1,,3,4", "T,,2,ספרייה,T001,0,30,,")).json()
    response = save(client, to_save(body))
    assert response.status_code == 200, response.json()
    assert response.json() == {"buildings_added": 1, "buildings_replaced": 0, "places_added": 2, "places_skipped": 0}
    lab = session.scalars(select(Place).where(Place.name == "T201")).one()
    assert (lab.institution_id, lab.capacity, count(session, Seat, place_id=lab.id)) == (braude.id, 12, 12)
    assert lab.opening_hours  # the campus's usual hours, like the form


def test_an_excel_file_with_english_headers_in_another_order(client, admin):
    client.user = admin
    data = xlsx([
        ["Place name", "Kind", "Floor", "Capacity", "Code", "Floors"],
        ["T101", "group room", 0, 6, "t", 2],
        [None, None, None, None, None, None],
        ["T102", "Group_Room", 1.0, "4", "T", "2"],
    ])
    body = preview(client, data).json()
    assert body["problems"] == []
    [t] = body["buildings"]
    assert t["code"] == "T"
    assert [(p["line"], p["name"], p["floor"], p["capacity"]) for p in t["places"]] == [(2, "T101", 0, 6), (4, "T102", 1, 4)]


@pytest.mark.parametrize(
    ("data", "encoding"),
    [(b"\xef\xbb\xbf", "utf-8"), (b"", "cp1255")],  # Excel's "CSV UTF-8", and plain CSV on a Hebrew Windows
)
def test_csv_as_excel_saves_it_in_hebrew(client, admin, data, encoding):
    client.user = admin
    body = preview(client, data + csv("T,בניין הניסיון,1,ספרייה,ספריה ראשית,0,20,,", encoding=encoding)).json()
    assert body["problems"] == []
    assert body["buildings"][0]["name"] == "בניין הניסיון"
    assert body["buildings"][0]["places"][0]["name"] == "ספריה ראשית"


def test_semicolons_too(client, admin):
    client.user = admin
    body = preview(client, csv("T;;1;ספרייה;L1;0;20;;", header=HEADER.replace(",", ";"))).json()
    assert [p["name"] for p in body["buildings"][0]["places"]] == ["L1"]


@pytest.mark.parametrize(
    ("line", "code"),
    [
        ("T,,2,ספרייה,L1,5,20,,", "floor_not_in_building"),
        ("T,,2,מטבח,L1,0,20,,", "row_kind"),
        ("T,,2,ספרייה,L1,0,,,", "row_capacity"),
        ("T,,2,חוות מחשבים,L1,0,,3,", "row_lab_size"),
        ("T,,2,ספרייה,,0,20,,", "row_place_name"),
        ("T,,2,ספרייה,L1,קומה,20,,", "row_floor"),
        ("T,,0,,,,,,", "row_floors"),
        ("ט!,,2,,,,,,", "row_building_code"),
        ("T,=HYPERLINK(\"x\"),2,,,,,,", "row_formula"),
        ("T,,2,ספרייה,@SUM(A1),0,20,,", "row_formula"),
    ],
)
def test_a_bad_row_is_named_with_its_line_and_left_out(client, admin, line, code):
    client.user = admin
    body = preview(client, csv("K,,1,ספרייה,K1,0,10,,", line)).json()
    assert body["problems"] == [{"line": 3, "code": code}]
    assert [b["code"] for b in body["buildings"]] == ["K"]


def test_the_same_place_twice_and_a_building_that_differs(client, admin):
    client.user = admin
    body = preview(client, csv(
        "T,א,2,ספרייה,L1,0,20,,", "T,,2,ספרייה,L1,1,20,,", "T,,3,ספרייה,L2,0,20,,", "T,שם אחר,2,ספרייה,L3,0,20,,",
    )).json()
    assert body["problems"] == [
        {"line": 3, "code": "row_duplicate"}, {"line": 4, "code": "row_building_differs"},
        {"line": 5, "code": "row_building_differs"},
    ]


@pytest.mark.parametrize(
    ("data", "code"),
    [
        (b"\x00\x01binary", "import_not_table"),
        (csv("T,,1,,,,,,", header="שם,קומות"), "import_missing_column"),
        (csv(), "import_empty"),
        (csv(*[f"T,,1,ספרייה,L{n},0,5,," for n in range(301)]), "import_too_many_rows"),
        (b"PK\x03\x04 broken", "import_not_xlsx"),
    ],
)
def test_a_file_that_is_not_such_a_table(client, admin, data, code):
    client.user = admin
    response = preview(client, data)
    assert (response.status_code, response.json()["detail"]) == (422, code)


def test_a_huge_or_bomb_file_is_refused_before_it_is_read(client, admin):
    client.user = admin
    assert preview(client, b"a" * (300 * 1024)).json()["detail"] == "import_too_big"
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"0" * (20 * 1024 * 1024))
    assert preview(client, bomb.getvalue()).json()["detail"] == "import_too_big"


# --- Who may --------------------------------------------------------------------------------


def test_a_student_and_another_institutions_admin_may_not(client, session, admin):
    data = csv("T,,1,,,,,,")
    assert preview(client, data).status_code == 403
    assert save(client, [{"code": "T", "floors_count": 1, "places": []}]).status_code == 403
    other = seed_demo(session)
    client.user = User(institution_id=other.id, email="x@y.example", display_name="X", role=UserRole.INSTITUTION_ADMIN)
    session.add(client.user)
    session.flush()
    assert preview(client, data).status_code == 404
    assert save(client, [{"code": "T", "floors_count": 1, "places": []}]).status_code == 404


def test_the_saved_table_is_checked_again_not_trusted(client, admin):
    client.user = admin
    bad = [{"code": "T", "floors_count": 1, "places": [{"kind": "library", "name": "L", "floor": 3, "capacity": 5}]}]
    assert save(client, bad).json()["detail"] == "floor_not_in_building"
    too_many = [{"code": f"T{n}", "floors_count": 1, "places": []} for n in range(101)]
    assert save(client, too_many).status_code == 422


# --- A building that is already there --------------------------------------------------------


def test_an_existing_building_is_named_with_what_is_in_it(client, admin, session, braude, student):
    m = next(b for b in braude.buildings if b.code == "M")
    lab = next(p for p in m.places if p.name == "M206")
    seat = session.scalars(select(Seat).where(Seat.place_id == lab.id)).first()
    session.add(Booking(
        institution_id=braude.id, user_id=student.id, place_id=lab.id, seat_id=seat.id,
        starts_at=datetime(2026, 10, 12, 9, tzinfo=UTC), ends_at=datetime(2026, 10, 12, 10, tzinfo=UTC),
        status=BookingStatus.BOOKED, source=BookingSource.ADVANCE,
    ))
    session.add(ClassSlot(user_id=student.id, weekday=0, starts=time(9), ends=time(10), building_id=m.id))
    session.flush()
    client.user = admin
    # The file says 5 floors, the building has 3: floor 4 fits the file only.
    body = preview(client, csv("m,,5,חוות מחשבים,M206,1,,2,2", "M,,5,ספרייה,M001,0,30,,", "M,,5,ספרייה,M002,4,30,,")).json()
    assert body["problems"] == []
    [b] = body["buildings"]
    assert b["existing"] == {"name": m.name, "floors_count": 3, "places": 2, "upcoming_bookings": 1, "classes": 1, "locked": False,
        "status": "active",
    }
    assert [(p["name"], p["when_added"]) for p in b["places"]] == [
        ("M206", "place_name_taken"), ("M001", None), ("M002", "floor_not_in_building"),
    ]


def test_add_puts_the_new_places_into_it_and_skips_the_ones_it_has(client, admin, session, braude):
    client.user = admin
    body = preview(client, csv("M,,3,חוות מחשבים,M206,1,,2,2", "M,,3,ספרייה,M001,0,30,,")).json()
    items = to_save(body)
    items[0]["places"].append({"kind": "computer_lab", "name": "M206", "floor": 1, "lab_rows": 2, "lab_cols": 2})
    response = save(client, items)
    assert response.json() == {"buildings_added": 0, "buildings_replaced": 0, "places_added": 1, "places_skipped": 1}
    m = session.scalars(select(Building).where(Building.institution_id == braude.id, Building.code == "M")).one()
    assert sorted(p.name for p in m.places) == ["M001", "M206", "M305"]
    assert count(session, Seat, place_id=next(p.id for p in m.places if p.name == "M206")) == 40  # untouched


def test_replace_deletes_it_and_adds_it_again_from_the_file(client, admin, session, braude, student):
    m = next(b for b in braude.buildings if b.code == "M")
    lab = next(p for p in m.places if p.name == "M206")
    seat = session.scalars(select(Seat).where(Seat.place_id == lab.id)).first()
    session.add(Booking(
        institution_id=braude.id, user_id=student.id, place_id=lab.id, seat_id=seat.id,
        starts_at=datetime(2026, 10, 12, 9, tzinfo=UTC), ends_at=datetime(2026, 10, 12, 10, tzinfo=UTC),
        status=BookingStatus.BOOKED, source=BookingSource.ADVANCE,
    ))
    session.add_all([
        ClassSlot(user_id=student.id, weekday=0, starts=time(9), ends=time(10), building_id=m.id),
        ClassSlot(user_id=student.id, weekday=1, starts=time(9), ends=time(10), building_id=m.id, place_text="m"),
    ])
    session.commit()
    client.user = admin
    body = preview(client, csv("M,בניין M,2,ספרייה,M001,0,30,,")).json()
    response = save(client, to_save(body, "replace"))
    assert response.json() == {"buildings_added": 0, "buildings_replaced": 1, "places_added": 1, "places_skipped": 0}
    session.expire_all()
    new = session.scalars(select(Building).where(Building.institution_id == braude.id, Building.code == "M")).one()
    assert (new.id != m.id, new.floors_count, [p.name for p in new.places]) == (True, 2, ["M001"])
    assert count(session, Booking) == 0
    # The students keep their classes, now in the new building of that name.
    classes = session.scalars(select(ClassSlot).order_by(ClassSlot.weekday)).all()
    assert [c.building_id for c in classes] == [new.id, new.id]


def test_replace_refuses_when_the_building_changed_since_the_preview(client, admin, session, braude):
    client.user = admin
    body = preview(client, csv("M,,2,ספרייה,M001,0,30,,")).json()
    items = to_save(body, "replace")
    items[0]["expected_places"] = 5
    assert save(client, items).json()["detail"] == "import_stale"
    del items[0]["expected_places"]
    assert save(client, items).status_code == 422  # a delete never goes without the number seen
    assert count(session, Building, code="M") == 1


def test_one_refusal_and_nothing_is_saved(client, admin, session, braude):
    client.user = admin
    before = count(session, Place)
    body = preview(client, csv("M,,2,ספרייה,M001,0,30,,", "T,,1,ספרייה,T1,0,5,,")).json()
    items = to_save(body, "replace")
    items[1]["places"].append({"kind": "library", "name": "T1", "floor": 0, "capacity": 5})  # twice
    assert save(client, items).json()["detail"] == "place_name_taken"
    session.expire_all()
    assert (count(session, Place), count(session, Building, code="T")) == (before, 0)
    assert {p.name for p in next(b for b in braude.buildings if b.code == "M").places} == {"M206", "M305"}


# --- The shared demo campus -----------------------------------------------------------------


def test_the_demo_campus_keeps_its_cap_and_its_seed_buildings(client, session, demo, demo_admin, demo_on):
    session.commit()
    client.user = demo_admin
    seed_code = demo.buildings[0].code
    items = [{
        "code": seed_code, "floors_count": 1, "action": "replace", "expected_places": len(demo.buildings[0].places),
        "expected_upcoming_bookings": 0, "places": [],
    }]
    assert save(client, items, "demo").json()["detail"] == "demo_campus_locked"
    many = [{"code": f"V{n}", "floors_count": 1, "places": []} for n in range(10)]
    assert save(client, many, "demo").json()["detail"] == "demo_campus_full"
    # A building a visitor added can be put back from the file.
    assert save(client, [{"code": "V1", "floors_count": 1, "places": []}], "demo").status_code == 200
    again = [{"code": "V1", "floors_count": 2, "action": "replace", "expected_places": 0, "expected_upcoming_bookings": 0, "places": []}]
    assert save(client, again, "demo").json()["buildings_replaced"] == 1


def test_a_booking_in_the_past_is_not_counted_as_upcoming(client, admin, session, braude, student):
    m = next(b for b in braude.buildings if b.code == "M")
    lab = next(p for p in m.places if p.name == "M206")
    seat = session.scalars(select(Seat).where(Seat.place_id == lab.id)).first()
    start = client.clock.now - timedelta(days=1)
    session.add(Booking(
        institution_id=braude.id, user_id=student.id, place_id=lab.id, seat_id=seat.id,
        starts_at=start, ends_at=start + timedelta(hours=1), status=BookingStatus.BOOKED, source=BookingSource.ADVANCE,
    ))
    session.flush()
    client.user = admin
    assert preview(client, csv("M,,3,,,,,,")).json()["buildings"][0]["existing"]["upcoming_bookings"] == 0


# --- The templates the admin page offers -------------------------------------------------

TEMPLATES = Path(__file__).resolve().parents[2] / "frontend" / "public" / "templates"


def test_the_templates_are_empty_and_ready():
    for name in ("buildings.xlsx", "buildings.csv"):
        with pytest.raises(campus_import.ImportRefused) as refused:
            campus_import.read_table((TEMPLATES / name).read_bytes())
        assert refused.value.code == "import_empty", name


def filled_template(rows: list[list]) -> bytes:
    """The Excel template as an admin fills it: from its first empty row."""
    book = load_workbook(TEMPLATES / "buildings.xlsx")
    sheet = book.worksheets[0]
    first = campus_import.TEMPLATE_FIRST_ROW
    for offset, row in enumerate(rows):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=first + offset, column=column, value=value)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_the_excel_template_filled_in_reads_whole():
    # short name, name, floors, status, kind, place, floor, people, rows, per row, Sun-Thu, Fri, note, lat, lng
    table = campus_import.read_table(filled_template([
        ["M", "בניין ההנדסה", 3, "פעיל", "חוות מחשבים", "M206", 1, None, 5, 8, "07:00-20:00", "07:00-14:00", "ליד המעלית"],
        [None, None, None, None, "חדר לימוד קבוצתי", "M310", 2, 6, None, None, None, "סגור"],
        ["NX", "בניין חדש", 2, "בבנייה", "אין מקומות לימוד"],
    ]))
    assert table.problems == []
    m, nx = table.buildings
    assert [(line, p.name) for line, p in m.places] == [(3, "M206"), (4, "M310")]
    assert (m.places[0][1].weekday_hours, m.places[0][1].location_note) == ("07:00-20:00", "ליד המעלית")
    assert (m.places[1][1].weekday_hours, m.places[1][1].friday_hours) == (None, "closed")
    assert (nx.code, nx.status, nx.places) == ("NX", BuildingStatus.UNDER_CONSTRUCTION, [])


def test_the_excel_templates_lists_and_headers_are_what_the_code_reads():
    sheet = load_workbook(TEMPLATES / "buildings.xlsx").worksheets[0]
    headers = [c.value for c in sheet[campus_import.TEMPLATE_FIRST_ROW - 1] if c.value]
    assert all(campus_import._key(h) in campus_import._HEADERS for h in headers), headers
    choices = [
        value.strip('"')
        for rule in sheet.data_validations.dataValidation
        for value in rule.formula1.strip('"').split(",")
    ]
    assert choices
    for choice in choices:
        key = campus_import._key(choice)
        assert key in campus_import._KINDS or key in campus_import._NO_PLACES or key in campus_import._STATUSES, choice


def test_an_empty_short_name_is_the_building_above(client, admin):
    client.user = admin
    body = preview(client, csv("T,בניין,2,ספרייה,L1,0,20,,", ",,,ספרייה,L2,1,20,,")).json()
    assert body["problems"] == []
    assert [p["name"] for p in body["buildings"][0]["places"]] == ["L1", "L2"]
    first = preview(client, csv(",,2,ספרייה,L1,0,20,,")).json()
    assert first["problems"] == [{"line": 2, "code": "row_building_code"}]


def test_no_study_places_adds_the_building_only(client, admin):
    client.user = admin
    body = preview(client, csv("T,,2,אין מקומות לימוד,,,,,")).json()
    assert (body["problems"], body["buildings"][0]["places"]) == ([], [])


HOURS_HEADER = HEADER + ",ימים א'-ה',יום ו',מצב הבניין"


@pytest.mark.parametrize(
    ("week", "friday", "expected"),
    [
        ("07:00-20:00", "07:00-14:00", ("07:00-20:00", "07:00-14:00")),
        ("7:00 - 20:00", "סגור", ("07:00-20:00", "closed")),
        ("8:30–16:00", "", ("08:30-16:00", None)),
    ],
)
def test_opening_hours_in_their_forms(client, admin, week, friday, expected):
    client.user = admin
    body = preview(client, csv(f"T,,1,ספרייה,L1,0,20,,,{week},{friday},", header=HOURS_HEADER)).json()
    assert body["problems"] == []
    place = body["buildings"][0]["places"][0]
    assert (place["weekday_hours"], place["friday_hours"]) == expected


@pytest.mark.parametrize(("cells", "code"), [("20:00-07:00,,", "row_hours"), ("כל היום,,", "row_hours"), (",,ישן", "row_status")])
def test_bad_hours_or_status(client, admin, cells, code):
    client.user = admin
    body = preview(client, csv(f"T,,1,ספרייה,L1,0,20,,,{cells}", header=HOURS_HEADER)).json()
    assert body["problems"] == [{"line": 2, "code": code}]


def test_the_hours_and_status_are_saved(client, admin, session, braude):
    client.user = admin
    data = csv("T,,1,ספרייה,L1,0,20,,,08:00-18:00,סגור,חדש", "T,,1,ספרייה,L2,0,20,,,,,", header=HOURS_HEADER)
    assert save(client, to_save(preview(client, data).json())).status_code == 200
    session.expire_all()
    t = session.scalars(select(Building).where(Building.institution_id == braude.id, Building.code == "T")).one()
    assert t.status == BuildingStatus.NEW
    hours = {p.name: {h.weekday: (h.opens, h.closes) for h in p.opening_hours} for p in t.places}
    assert hours["L1"] == {d: (time(8), time(18)) for d in (6, 0, 1, 2, 3)}  # Friday closed, Saturday too
    assert hours["L2"][4] == (time(7), time(14))  # the campus's usual hours


def test_the_saved_hours_are_checked_again(client, admin):
    client.user = admin
    place = {"kind": "library", "name": "L", "floor": 0, "capacity": 5, "weekday_hours": "20:00-07:00"}
    assert save(client, [{"code": "T", "floors_count": 1, "places": [place]}]).status_code == 422


# --- After the security review -------------------------------------------------------------


@pytest.mark.parametrize("hidden", ["\u200b", "\ufeff", "\u2060", "\x01"])
def test_a_hidden_character_does_not_slip_a_formula_past(client, admin, hidden):
    client.user = admin
    body = preview(client, csv(f"T,{hidden}=HYPERLINK(\"x\"),1,,,,,,", f"Q,,1,ספרייה,{hidden}=1+1,0,5,,")).json()
    assert body["problems"] == [{"line": 2, "code": "row_formula"}, {"line": 3, "code": "row_formula"}]


@pytest.mark.parametrize(
    "where",
    [{"name": "=HYPERLINK(1)"}, {"places": [{"kind": "library", "name": "@SUM(1)", "floor": 0, "capacity": 5}]},
     {"places": [{"kind": "library", "name": "L", "floor": 0, "capacity": 5, "location_note": "+cmd"}]}],
)
def test_the_saved_table_and_the_forms_refuse_formulas_too(client, admin, session, where):
    client.user = admin
    assert save(client, [{"code": "T", "floors_count": 1, "places": [], **where}]).status_code == 422
    if "name" in where:
        assert client.post("/admin/institutions/braude/buildings", json={"code": "T", "floors_count": 1, **where}).status_code == 422
    assert client.post("/admin/institutions/braude/buildings", json={"code": "-1", "floors_count": 1}).status_code == 422


def test_replace_refuses_when_a_booking_was_made_since_the_preview(client, admin, session, braude, student):
    client.user = admin
    body = preview(client, csv("M,,2,ספרייה,M001,0,30,,")).json()
    items = to_save(body, "replace")
    assert items[0]["expected_upcoming_bookings"] == 0
    m = next(b for b in braude.buildings if b.code == "M")
    lab = next(p for p in m.places if p.name == "M206")
    seat = session.scalars(select(Seat).where(Seat.place_id == lab.id)).first()
    start = client.clock.now + timedelta(hours=2)
    session.add(Booking(
        institution_id=braude.id, user_id=student.id, place_id=lab.id, seat_id=seat.id,
        starts_at=start, ends_at=start + timedelta(hours=1), status=BookingStatus.BOOKED, source=BookingSource.ADVANCE,
    ))
    session.commit()
    assert save(client, items).json()["detail"] == "import_stale"
    assert count(session, Booking) == 1


def test_blank_lines_do_not_hide_rows_past_a_cut(client, admin):
    client.user = admin
    lines = [""] * 30 + [f"T,,1,ספרייה,L{n},0,5,," for n in range(300)]
    body = preview(client, csv(*lines)).json()
    assert len(body["buildings"][0]["places"]) == 300
    many = [""] * 2000 + ["T,,1,,,,,,"]
    assert preview(client, csv(*many)).json()["detail"] == "import_too_many_rows"


def test_a_table_of_more_buildings_than_one_save_takes_is_refused_at_once(client, admin):
    client.user = admin
    response = preview(client, csv(*[f"B{n},,1,,,,,," for n in range(101)]))
    assert response.json()["detail"] == "import_too_many_buildings"


def test_one_import_creates_a_bounded_number_of_stations(client, admin):
    client.user = admin
    labs = [{"kind": "computer_lab", "name": f"L{n}", "floor": 0, "lab_rows": 20, "lab_cols": 20} for n in range(6)]
    assert save(client, [{"code": "T", "floors_count": 1, "places": labs}]).status_code == 422
    assert save(client, [{"code": "T", "floors_count": 1, "places": labs[:5]}]).status_code == 200


# --- After the code review -----------------------------------------------------------------


def test_semicolons_under_a_title_row(client, admin):
    client.user = admin
    data = "רשימת בניינים\n".encode() + csv("T;;1;ספרייה;L1;0;20;;", header=HEADER.replace(",", ";"))
    assert [p["name"] for p in preview(client, data).json()["buildings"][0]["places"]] == ["L1"]


def test_a_bad_first_row_does_not_get_the_next_ones_blamed(client, admin):
    client.user = admin
    body = preview(client, csv("T,א,2,מטבח,L1,0,20,,", "T,,,ספרייה,L2,0,20,,")).json()
    assert body["problems"] == [{"line": 2, "code": "row_kind"}]
    assert [p["name"] for p in body["buildings"][0]["places"]] == ["L2"]


def test_a_building_name_that_is_a_number(client, admin):
    client.user = admin
    data = xlsx([["Code", "Building name", "Floors"], ["T", 2020, 1]])
    assert preview(client, data).json()["buildings"][0]["name"] == "2020"


def test_a_building_new_at_the_preview_and_added_elsewhere_since(client, admin, session):
    client.user = admin
    items = to_save(preview(client, csv("T,,1,ספרייה,L1,0,5,,")).json())
    assert client.post("/admin/institutions/braude/buildings", json={"code": "T", "floors_count": 3}).status_code == 201
    assert save(client, items).json()["detail"] == "import_stale"


def test_replace_keeps_classes_whatever_words_they_were_typed_in(client, admin, session, braude, student):
    m = next(b for b in braude.buildings if b.code == "M")
    session.add(ClassSlot(user_id=student.id, weekday=0, starts=time(9), ends=time(10), building_id=m.id, place_text="בניין מיטשל"))
    session.commit()
    client.user = admin
    save(client, to_save(preview(client, csv("M,,2,,,,,,")).json(), "replace"))
    session.expire_all()
    new = session.scalars(select(Building).where(Building.institution_id == braude.id, Building.code == "M")).one()
    assert session.scalars(select(ClassSlot)).one().building_id == new.id


def test_the_preview_says_a_demo_seed_building_cannot_be_replaced(client, session, demo, demo_admin, demo_on):
    session.commit()
    client.user = demo_admin
    code = demo.buildings[0].code
    body = preview(client, csv(f"{code},,1,,,,,,", "NEW,,1,,,,,,"), "demo").json()
    assert [b["existing"] and b["existing"]["locked"] for b in body["buildings"]] == [True, None]


# --- After the review of the template -----------------------------------------------------


def test_a_missing_short_name_on_a_new_building_is_not_taken_for_the_one_above(client, admin):
    client.user = admin
    body = preview(client, csv("M,בניין M,3,ספרייה,M1,0,20,,", ",בניין B,2,ספרייה,B1,0,20,,", ",,,ספרייה,B2,1,20,,")).json()
    assert body["problems"] == [{"line": 3, "code": "row_building_code"}, {"line": 4, "code": "row_building_above"}]
    assert [p["name"] for p in body["buildings"][0]["places"]] == ["M1"]


def test_rows_under_a_bad_building_row_say_so(client, admin):
    client.user = admin
    body = preview(client, csv("T,,0,ספרייה,L1,0,20,,", ",,,ספרייה,L2,0,20,,")).json()
    assert body["problems"] == [{"line": 2, "code": "row_floors"}, {"line": 3, "code": "row_building_above"}]


@pytest.mark.parametrize("header", ["ימים א׳-ה׳,יום ו׳", "ימים א’-ה’,יום ו’"])
def test_hours_headers_with_a_hebrew_geresh(client, admin, header):
    client.user = admin
    body = preview(client, csv("T,,1,ספרייה,L1,0,20,,,סגור,08:00-12:00", header=HEADER + "," + header)).json()
    place = body["buildings"][0]["places"][0]
    assert (place["weekday_hours"], place["friday_hours"]) == ("closed", "08:00-12:00")


def test_a_place_closed_every_day_is_refused(client, admin):
    client.user = admin
    body = preview(client, csv("T,,1,ספרייה,L1,0,20,,,סגור,סגור,", header=HOURS_HEADER)).json()
    assert body["problems"] == [{"line": 2, "code": "row_hours"}]
    place = {"kind": "library", "name": "L", "floor": 0, "capacity": 5, "weekday_hours": "closed", "friday_hours": "closed"}
    assert save(client, [{"code": "T", "floors_count": 1, "places": [place]}]).status_code == 422


def test_no_study_places_with_place_details_is_a_problem(client, admin):
    client.user = admin
    body = preview(client, csv("T,,1,אין מקומות לימוד,A,0,10,,")).json()
    assert body["problems"] == [{"line": 2, "code": "row_no_places_details"}]


def test_the_preview_shows_the_status_of_a_building_that_is_there(client, admin, session, braude):
    m = next(b for b in braude.buildings if b.code == "M")
    m.status = BuildingStatus.UNDER_CONSTRUCTION
    session.commit()
    client.user = admin
    assert preview(client, csv("M,,3,,,,,,")).json()["buildings"][0]["existing"]["status"] == "under_construction"
