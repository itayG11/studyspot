"""Builds the empty templates the admin page offers for a table of buildings
and places: frontend/public/templates/buildings.xlsx and buildings.csv.

    cd backend && .venv/bin/python scripts/make_campus_template.py

Styled like the campus data sheet of stage 1 (docs/CAMPUS_DATA.md): a row
of groups over the headers, a list to choose the kind of place from, a
colour for each kind, grey on the columns that kind does not use, and a
sheet of instructions. The headers and the lists come from
app/campus_import.py, and tests/test_campus_import.py checks that the file
reads with it."""

import csv
import io
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.campus_import import COLUMNS, KINDS, NO_PLACES, STATUSES, TEMPLATE_FIRST_ROW
from app.models import BuildingStatus

OUT = Path(__file__).resolve().parents[2] / "frontend" / "public" / "templates"
ROWS = 300  # app/campus_import.py MAX_ROWS

# (field, width), in the order of the sheet; each header is the field's
# first name in COLUMNS.
FIELDS = [
    ("code", 10), ("name", 22), ("floors_count", 11), ("status", 12),
    ("kind", 20), ("place", 16), ("floor", 8), ("capacity", 11),
    ("lab_rows", 9), ("lab_cols", 12),
    ("weekday_hours", 13), ("friday_hours", 12),
    ("location_note", 26),
    ("latitude", 11), ("longitude", 11),
]
GROUPS = [  # (title, first field, last field, colour)
    ("הבניין", "code", "status", "1F4E79"),
    ("המקום", "kind", "capacity", "7030A0"),
    ("חוות מחשבים", "lab_rows", "lab_cols", "375623"),
    ("שעות פתיחה", "weekday_hours", "friday_hours", "C55A11"),
    ("איך מגיעים", "location_note", "location_note", "7030A0"),
    ("מיקום על המפה, לא חובה", "latitude", "longitude", "595959"),
]
KIND_COLOURS = {"group_room": "A9D08E", "open_area": "F4B183", "library": "E89AB8", "computer_lab": "8FD3D9"}
GREY = PatternFill("solid", fgColor="A6A6A6")

INSTRUCTIONS = [
    "איך ממלאים",
    "",
    "1. כל בניין הוא בלוק של שורות. בשורה הראשונה שלו כותבים שם קצר, שם, מספר קומות ומצב.",
    "2. כל מקום לימוד בבניין מקבל שורה משלו. בשורות שמתחת משאירים את פרטי הבניין ריקים, וזה אומר \"אותו בניין כמו למעלה\".",
    "3. השם הקצר: אותיות באנגלית, ספרות ומקף, כמו M או EM. הוא מופיע בשמות המקומות ועל השלטים.",
    "4. בעמודה \"סוג המקום\" בוחרים מהרשימה. העמודות שלא שייכות לסוג נצבעות באפור, ואותן משאירים ריקות.",
    "5. בניין בלי מקומות לימוד: שורה אחת, ובסוג המקום בוחרים \"אין מקומות לימוד\". הוא יופיע על המפה.",
    "6. קומה: הכניסה היא קומה 0, ומעליה 1, 2 וכן הלאה.",
    "7. חוות מחשבים: שורות ועמדות בשורה. נוצרת עמדה לכל שורה ועמודה, ומהן מפת העמדות.",
    "8. שעות פתיחה: למשל 07:00-20:00, או \"סגור\". ריק נותן את השעות הרגילות: א'-ה' 07:00-20:00, ו' 07:00-14:00. שבת סגורה.",
    "9. מיקום על המפה לא חובה: אפשר לקבוע אותו אחר כך בדף הניהול, ב\"מיקום בניינים\".",
    "10. כשמעלים, האתר מראה קודם מה נמצא בטבלה, ושורה עם בעיה מסומנת עם הסיבה. שום דבר לא נשמר עד שמאשרים.",
    "",
    "דוגמה ממולאת יש בגיליון \"דוגמה\". ממלאים את הגיליון \"בניינים ומקומות\", והוא היחיד שנקרא.",
]
EXAMPLE = [
    ["M", "בניין ההנדסה", 3, "פעיל", "חוות מחשבים", "M206", 1, None, 5, 8, "07:00-20:00", "07:00-14:00", "ליד המעלית"],
    [None, None, None, None, "חדר לימוד קבוצתי", "M310", 2, 6, None, None, None, "סגור"],
    ["NX", "בניין חדש", 2, "בבנייה", "אין מקומות לימוד"],
]


def header(field: str) -> str:
    return COLUMNS[field][0]


def position(field: str) -> int:
    return [f for f, _ in FIELDS].index(field) + 1


def column(field: str) -> str:
    return get_column_letter(position(field))


def workbook() -> Workbook:
    book = Workbook()
    sheet = book.active
    sheet.title = "בניינים ומקומות"
    sheet.sheet_view.rightToLeft = True
    white = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="BFBFBF")
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for title, first, last, colour in GROUPS:
        if first != last:
            sheet.merge_cells(f"{column(first)}1:{column(last)}1")
        cell = sheet[f"{column(first)}1"]
        cell.value, cell.font, cell.alignment = title, white, centre
        for index in range(position(first), position(last) + 1):
            sheet.cell(row=1, column=index).fill = PatternFill("solid", fgColor=colour)
    heads = TEMPLATE_FIRST_ROW - 1
    for index, (field, width) in enumerate(FIELDS, start=1):
        cell = sheet.cell(row=heads, column=index, value=header(field))
        cell.font, cell.alignment = white, centre
        cell.fill = PatternFill("solid", fgColor="2E75B6")
        cell.border = Border(bottom=thin)
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[heads].height = 36
    sheet.freeze_panes = f"B{TEMPLATE_FIRST_ROW}"

    last = TEMPLATE_FIRST_ROW + ROWS - 1
    kind = f"${column('kind')}{TEMPLATE_FIRST_ROW}"
    kinds = [aliases[0] for aliases in KINDS.values()] + [NO_PLACES[0]]
    for field, values in (("kind", kinds), ("status", [STATUSES[s][0] for s in BuildingStatus])):
        rule = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        rule.error, rule.errorTitle, rule.showErrorMessage = "בוחרים ערך מהרשימה.", "ערך לא מוכר", True
        rule.add(f"{column(field)}{TEMPLATE_FIRST_ROW}:{column(field)}{last}")
        sheet.add_data_validation(rule)

    def grey(first: str, last_field: str, when: str) -> None:
        span = f"{column(first)}{TEMPLATE_FIRST_ROW}:{column(last_field)}{last}"
        sheet.conditional_formatting.add(span, FormulaRule(formula=[f'{kind}="{when}"'], fill=GREY, stopIfTrue=False))

    for name, colour in KIND_COLOURS.items():
        label = KINDS[name][0]
        span = f"{column('kind')}{TEMPLATE_FIRST_ROW}:{column('kind')}{last}"
        sheet.conditional_formatting.add(
            span, FormulaRule(formula=[f'{kind}="{label}"'], fill=PatternFill("solid", fgColor=colour))
        )
        if name == "computer_lab":
            grey("capacity", "capacity", label)
        else:
            grey("lab_rows", "lab_cols", label)
    grey("place", "location_note", NO_PLACES[0])

    guide = book.create_sheet("הוראות")
    guide.sheet_view.rightToLeft = True
    guide.column_dimensions["A"].width = 115
    for index, line in enumerate(INSTRUCTIONS, start=1):
        guide.cell(row=index, column=1, value=line).alignment = Alignment(wrap_text=True)
    guide["A1"].font = Font(bold=True, size=14)

    # The example on a sheet of its own, so that the one to fill stays empty.
    example = book.create_sheet("דוגמה")
    example.sheet_view.rightToLeft = True
    for index, (field, width) in enumerate(FIELDS, start=1):
        example.cell(row=1, column=index, value=header(field)).font = Font(bold=True)
        example.column_dimensions[get_column_letter(index)].width = width
    for row, values in enumerate(EXAMPLE, start=2):
        for offset, value in enumerate(values, start=1):
            example.cell(row=row, column=offset, value=value)
    return book


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    workbook().save(OUT / "buildings.xlsx")
    text = io.StringIO()
    csv.writer(text, lineterminator="\r\n").writerow([header(f) for f, _ in FIELDS])
    # With a BOM, so that Excel opens the Hebrew right.
    (OUT / "buildings.csv").write_bytes(b"\xef\xbb\xbf" + text.getvalue().encode("utf-8"))


if __name__ == "__main__":
    main()
